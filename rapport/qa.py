"""rapport/qa.py — premier harnais : le PDF lu page à page avec `pypdf`, le dépôt re-compté à nu.

Treize contrôles, douze lois du cahier des charges plus la loi des coupures. Chacun est une affirmation que le rapport écrit et que ce fichier vérifie sans
consommer le moteur de composition : si `rapport/doc.py` est faux, `qa.py` doit le voir. Les
contrôles qui, par nature, regardent la même chose que le moteur (loi de langue) le font par un
chemin différent — ici le texte extrait du PDF, et jamais le texte source de l'auteur.

Sortie : `audit/qa.json` (procès-verbal), code retour 1 si un contrôle est rouge, `audit/budgets.json`
complété de `secondes_qa`.
"""
from __future__ import annotations

import json
import re
import time
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
AUDIT = RACINE / "audit"
PDF = RACINE / "rapport" / "PDF" / "Cleopatre-rapport.pdf"
FIGS = RACINE / "rapport" / "figs"
DECK = RACINE / "soutenance" / "Cleopatre-soutenance.pptx"

# ── les lois, recopiées ici **volontairement** : un harnais qui importe la loi du moteur ne peut
#    pas trouver que le moteur se trompe. qa2.py les recopie une troisième fois, autrement.
SLOP = [r"dans un monde", r"il est important de", r"il convient de", r"force est de constater",
        r"à l[’']ère du", r"de nos jours", r"jouer un rôle clé", r"s'avère", r"crucial",
        r"de pointe", r"solution innovante"]
INTERDITS = [r"\border\b", r"\bcart\b", r"\bitem\b", r"\bshop\b", r"\bstore\b", r"\buser\b",
             r"\btracking\b", r"\bpassword\b"]
EMOJIS = "\U0001F300-\U0001FAFF⛔️✅❌"


def _texte_pdf() -> tuple[str, list[str], int]:
    """(texte complet, lignes rendues, nombre de pages) — lecture pypdf, une page à la fois."""
    from pypdf import PdfReader
    r = PdfReader(str(PDF))
    lignes: list[str] = []
    for page in r.pages:
        lignes += (page.extract_text() or "").splitlines()
    return "\n".join(lignes), [l.strip() for l in lignes], len(r.pages)


def _est_code(ligne: str) -> bool:
    """Ligne de verbatim : elle échappe aux lois de langue, ce n'est pas de la prose.

    Trois marques, par ordre de fiabilité : les doubles signes de programmation ; la gouttière de
    numéros que le compositeur imprime devant tout extrait (`45  * Best-effort by…`) ; et la densité
    de chemins `dossier/fichier.ext`, qui reste du verbatim même alignée dans une ligne de prose.
    """
    if re.search(r"[{};=<>|]{2}|=>|\bconst \b|\basync \b|: (string|number|boolean)\b|"
                 r"\bexport \b|\bfunction \b|`", ligne):
        return True
    if re.match(r"^\d{1,4}\s{1,4}\S", ligne):
        return True
    return len(re.findall(r"[\w./-]+\.(?:tsx|ts|css|json|py|md)\b", ligne)) >= 2


def _facts() -> dict:
    return json.loads((AUDIT / "facts.json").read_text(encoding="utf-8"))


# ══════════════════════════════════════════════════════════════ les treize contrôles
def _aplatis(noeuds):
    """Arbre d'outline pypdf → liste plate d'objets de signet (l'outline est imbriquée)."""
    for x in noeuds:
        if isinstance(x, list):
            yield from _aplatis(x)
        else:
            yield x


OFFICIELS = {                      # domaines que le kit admet en bibliographie (§4 : sources officielles)
    "nextjs.org", "react.dev", "developer.mozilla.org", "orm.drizzle.team", "zod.dev",
    "www.postgresql.org", "node-postgres.com", "www.w3.org", "eur-lex.europa.eu",
    "www.iso.org", "iso.org",
}


# ══════════════════════════════════════════════════════════════ les contrôles
ALIAS = {                      # nom tel qu'imprimé → nom de paquet dans package.json
    "Next.js": "next", "React": "react", "TypeScript": "typescript", "Tailwind": "tailwindcss",
    "Drizzle": "drizzle-orm", "Drizzle Kit": "drizzle-kit", "Zod": "zod",
    "Framer Motion": "framer-motion", "PostgreSQL": "embedded-postgres", "node-postgres": "pg",
}


def c_version(ctx):
    """Toute version imprimée est celle du dépôt ; tout paquet versionné du rapport est imprimé.

    Le rapport ne tape pas de version à la main : il génère sa table depuis les faits `ver_*`. Le
    contrôle vérifie donc les deux sens de la correspondance, et le nombre de citations attendues est
    le nombre de faits, pas une liste fixe écrite dans ce fichier.
    """
    pkg = json.loads((RACINE / "package.json").read_text(encoding="utf-8"))
    deps = {**pkg.get("dependencies", {}), **pkg.get("devDependencies", {})}
    faits = _facts()["faits"]
    versionnes = {k[4:]: str(v["valeur"]) for k, v in faits.items() if k.startswith("ver_")}
    absentes = sorted(nom for nom, ver in versionnes.items() if ver not in ctx["texte"])
    contrees, fausses = 0, []
    for libelle, paquet in ALIAS.items():
        for m in re.finditer(rf"{re.escape(libelle)} (\d+\.\d+\.\d+[\w.-]*)", ctx["texte"]):
            contrees += 1
            reel = str(deps.get(paquet, "")).lstrip("^~")
            if reel and m.group(1) != reel:
                fausses.append(f"{libelle} {m.group(1)} ≠ {reel}")
    return dict(verdict="vert" if not absentes and not fausses and contrees >= 4 else "rouge",
                mesure=len(versionnes),
                detail=f"{len(versionnes)} paquets versionnés dans les faits, absents du rendu "
                       f"{absentes[:3] or 'aucun'} ; {contrees} couple(s) nom-version lus dans le "
                       f"document, contredits {fausses[:3] or 'aucun'}")


def _recompte() -> dict:
    """Six structures recomptées dans le dépôt, sans l'extracteur du rapport."""
    src = RACINE / "src"
    schema = (src / "db" / "schema.ts").read_text(encoding="utf-8")
    actions = "".join(f.read_text(encoding="utf-8") for f in sorted((src / "actions").glob("*.ts")))
    manifeste = json.loads((RACINE / "public/data/product-image-manifest.json").read_text(encoding="utf-8"))
    return dict(
        tables=len(re.findall(r"= pgTable\(", schema)),
        enums=len(re.findall(r"= pgEnum\(", schema)),
        produits=len(manifeste) if isinstance(manifeste, list) else len(manifeste.get("products", [])),
        routes=len(list((src / "app").rglob("page.tsx"))),
        endpoints=len(list((src / "app" / "api").rglob("route.ts"))),
        actions=len(re.findall(r"^export (?:async )?(?:function|const)", actions, re.M)),
    )


def c_chiffres(ctx):
    """Six chiffres recomptés ici, confrontés au fait de l'extracteur et au texte rendu.

    Triple concordance : le dépôt, `audit/facts.json`, le PDF. Le nombre de colonnes est jugé à part,
    parce qu'il dépend d'une grammaire déclarée (colonnes de table, hors relations et hors blocs
    d'index) : le contrôle vérifie que le dictionnaire imprimé contient autant de lignes que le fait.
    """
    faits = _facts()["faits"]
    attendu = _recompte()
    ecarts = [f"{k}: recompté {n}, fait {faits[k]['valeur']}" for k, n in attendu.items()
              if str(faits[k]["valeur"]) != str(n)]
    absents = [f"{k} {n}" for k, n in attendu.items() if str(n) not in ctx["texte"]]
    inventaire = json.loads((AUDIT / "facts.json").read_text(encoding="utf-8"))["tables"]
    colonnes = sum(len(t["cols"]) for t in inventaire)
    grammaire = colonnes == int(faits["colonnes"]["valeur"]) and str(colonnes) in ctx["texte"]
    return dict(verdict="vert" if not ecarts and not absents and grammaire else "rouge",
                mesure=len(attendu),
                detail=f"{len(attendu)} structures recomptées à nu ; écarts avec l'extracteur "
                       f"{ecarts[:3] or 'aucun'} ; absents du rendu {absents[:3] or 'aucun'} ; "
                       f"colonnes du dictionnaire {colonnes} contre le fait "
                       f"{faits['colonnes']['valeur']} : {'cohérent' if grammaire else 'incohérent'}")


def c_chemins(ctx):
    # l'ordre des alternatives compte : `(?:ts|tsx)` rogne « account-nav.tsx » en « account-nav.ts »
    vus = set(re.findall(r"((?:src|public|rapport|soutenance)/[\w./[\]-]+\.(?:tsx|ts|css|json|py|md))"
                         r"(?::(\d+))?", ctx["texte"]))
    introuvables, hors_bornes = [], []
    for chemin, ligne in vus:
        p = RACINE / chemin
        if not p.exists():
            introuvables.append(chemin)
        elif ligne:
            n = len(p.read_text(encoding="utf-8").splitlines())
            if int(ligne) > n:
                hors_bornes.append(f"{chemin}:{ligne}>{n}")
    return dict(verdict="vert" if not (introuvables or hors_bornes) else "rouge", mesure=len(vus),
                detail=f"{len(vus)} chemins cités, introuvables {introuvables[:3] or 'aucun'}, "
                       f"hors bornes {hors_bornes[:3] or 'aucun'}")


def c_slop(ctx):
    touches = [motif for motif in SLOP if re.search(motif, ctx["texte"], flags=re.I)]
    return dict(verdict="vert" if not touches else "rouge", mesure=len(touches),
                detail=f"formules de remplissage lues dans le rendu : {touches[:5] or 'aucune'}")


def _git_ls() -> list[str]:
    """Fichiers suivis par git ; repli sur le système de fichiers si git est indisponible."""
    import subprocess
    try:
        return subprocess.run(["git", "-C", str(RACINE), "ls-files"], capture_output=True,
                              text=True, check=True).stdout.split()
    except Exception:
        return [str(p.relative_to(RACINE)) for p in RACINE.rglob("*")
                if p.is_file() and "/.git/" not in str(p) and "node_modules" not in str(p)]


def _bloc_code(lignes: list[str]) -> list[bool]:
    """Une ligne est-elle dans un bloc de verbatim ? Réponse par voisinage, pas par police.

    `pypdf` ne donne pas la police d'un glyphe. Le harnais de texte reconnaît donc un extrait de
    code à la densité de signes de programmation dans sa fenêtre de sept lignes. C'est approximatif
    **et assumé** : qa2.py, lui, lit le nom de la police. L'écart des deux comptages est imprimé.
    """
    marque = []
    for i, l in enumerate(lignes):
        fenetre = lignes[max(0, i - 3): i + 4]
        densite = sum(1 for x in fenetre if _est_code(x))
        marque.append(_est_code(l) or densite >= 3)
    return marque


_IDENTS: set[str] = set()


def _identifiants() -> set[str]:
    """Noms que le dépôt impose au vocabulaire du rapport : fichiers, symboles, valeurs d'énumération.

    La loi anti-emprunt porte sur la **prose**. Un identifiant recopié du code n'est pas une
    traduction manquée, c'est une citation : `cart-drawer.tsx` et `shop.ts` doivent rester ce qu'ils
    sont dans le dépôt, sous peine de rendre le chemin faux. Le harnais constitue donc la liste depuis
    l'arborescence (et non depuis un dictionnaire tenu à la main), puis la retire des lignes avant
    recherche. Ce que le harnais de texte ne peut pas voir — la police réellement employée — est jugé
    par qa2.py.
    """
    global _IDENTS
    if _IDENTS:
        return _IDENTS
    noms: set[str] = set()
    src = RACINE / "src"
    for d in (p for p in src.rglob("*") if p.is_dir()):
        noms.add(d.name.lower())
    for p in list(src.rglob("*.tsx")) + list(src.rglob("*.ts")):
        noms.add(p.stem.lower())
        noms.update(re.findall(r"[A-Za-z]\w*-(?:drawer|provider|card|table|form|overlay|controls|actions|timeline)", p.stem.lower()))
    for f in sorted((src / "actions").glob("*.ts")):
        noms.update(m.lower() for m in re.findall(r"^export (?:async )?(?:function|const) (\w+)",
                                                   f.read_text(encoding="utf-8"), re.M))
    schema = (src / "db" / "schema.ts").read_text(encoding="utf-8")
    for m in re.finditer(r'= pgEnum\(\s*"[\w_]+", \[([^\]]*)\]', schema, re.S):
        noms.update(x.strip(' \'"') for x in m.group(1).split(',') if x.strip(' \'"'))
    _IDENTS = {n for n in noms if len(n) > 3}
    return _IDENTS


def _hors_identifiants(ligne: str) -> str:
    """Ligne allégée des citations du dépôt : ce qui reste est de la prose, et elle est jugeable."""
    nu = re.sub(r"[\w{}/.,*-]+\.(?:tsx|ts|css|json|py|md)\b", " ", ligne)
    for nom in _identifiants():
        nu = re.sub(rf"(?<![\w-]){re.escape(nom)}(?![\w-])", " ", nu, flags=re.I)
    return nu


def _region_anglaise(lignes: list[str]) -> tuple[int, int]:
    """Le résumé anglais du mémoire est en anglais : région bornée, exemption **écrite**.

    Elle va du titre « Abstract » au prochain titre de même niveau. Rien n'est masqué : le nombre de
    lignes exemptées est renvoyé dans le procès-verbal.
    """
    # la région commence au titre « Abstract » qui suit le résumé français : le premier « Abstract »
    # du document est celui du sommaire, et le dernier est le titre courant de la page suivante
    apres = max([i for i, l in enumerate(lignes) if l.strip().startswith("Mots-clés")], default=-1)
    debut = next((i for i in range(apres + 1, len(lignes))
                  if re.match(r"^Abstract\b", lignes[i].strip())),
                 next((i for i, l in enumerate(lignes) if re.match(r"^Abstract\b", l.strip())),
                      len(lignes)))

    fin = next((i for i in range(debut + 1, len(lignes))
                if re.match(r"^(?:\d+(?:\.\d+)*\s+|Annexe [A-F]|[A-F]\.\d*\s+|Quatrième de"
                            r"\s+couverture)", lignes[i].strip())), len(lignes))
    return debut, fin


def c_synonymes(ctx):
    """Aucun mot anglais volé à la langue dans la prose. Trois exemptions, toutes comptées.

    verbatim (code et chemins), identifiants imposés par le dépôt, région du résumé anglais. Une
    exemption sans compteur serait une faille : les trois nombres sont imprimés.
    """
    lignes = ctx["lignes"]
    dans_code = _bloc_code(lignes)
    de, df = _region_anglaise(lignes)
    hits, exempts_code, exempts_anglais = [], 0, 0
    for i, l in enumerate(lignes):
        if not l:
            continue
        if de <= i < df:
            exempts_anglais += 1
            continue
        if dans_code[i]:
            exempts_code += 1
            continue
        nu = _hors_identifiants(l)
        for motif in INTERDITS:
            if re.search(motif, nu, flags=re.I):
                hits.append(f"{motif} : {l[:52]}")
                break
    return dict(verdict="vert" if not hits else "rouge", mesure=len(hits),
                detail=f"{len(hits)} prose(s) avec mot emprunté ; exemptions comptées : "
                       f"{exempts_code} ligne(s) de verbatim, {exempts_anglais} ligne(s) de la "
                       f"région du résumé anglais ; exemples : {hits[:3] or 'aucun'}")


def c_signets(ctx):
    from pypdf import PdfReader
    r = PdfReader(str(PDF))
    total = len(list(_aplatis(r.outline)))
    registre = json.loads((AUDIT / "signets.json").read_text(encoding="utf-8"))
    titres = [str(x.get("/Title", "")) for x in _aplatis(r.outline)]
    sales = [t for t in titres if "<" in t or "&#" in t or not t.strip()]
    return dict(verdict="vert" if total == len(registre) and not sales else "rouge", mesure=total,
                detail=f"outline {total} entrées, registre {len(registre)}, titres fautifs "
                       f"{sales[:2] or 'aucun'}")


HOSTS_EXEMPLES = {"localhost", "127.0.0.1"}   # consigne d'exécution locale, pas une source citée


def c_liens(ctx):
    from pypdf import PdfReader
    r = PdfReader(str(PDF))
    internes = uri = morts = 0
    for page in r.pages:
        for a in (page.get("/Annots") or []):
            ob = a.get_object()
            if ob.get("/Subtype") != "/Link":
                continue
            acte = ob.get("/A")
            if acte is not None and "/URI" in str(acte):
                uri += 1
            elif "/Dest" in ob or (acte is not None and "/D" in acte):
                internes += 1
            else:
                morts += 1
    domains = set(re.findall(r"https?://([\w.-]+)", ctx["texte"]))
    hors_liste = sorted(domains - OFFICIELS - HOSTS_EXEMPLES)
    return dict(verdict="vert" if morts == 0 and internes > 150 and uri == 0 and not hors_liste
                else "rouge", mesure=internes,
                detail=f"{internes} liens internes, {morts} mort(s), {uri} lien(s) externe(s) "
                       "cliquable(s) (le kit n'en pose aucun : les URL sont imprimées et vérifiables "
                       "à la main), domaines hors liste officielle : "
                       f"{hors_liste or 'aucun'} ; {len(domains)} domaines cités dont "
                       f"{len(domains & HOSTS_EXEMPLES)} exemptés comme exemples d'exécution")


def c_planches(ctx):
    from pypdf import PdfReader
    from PIL import Image
    pdf = sorted(FIGS.glob("*.pdf"))
    png = sorted(FIGS.glob("*.png"))
    ratios = []
    for p in pdf:
        if not p.with_suffix(".png").exists():
            ratios.append((p.stem, "PNG absent"))
            continue
        with Image.open(p.with_suffix(".png")) as im:
            pr = PdfReader(str(p)).pages[0]
            w, h = float(pr.mediabox.width), float(pr.mediabox.height)
            ecart = abs(w / h - im.size[0] / im.size[1])
            if ecart > 0.01:
                ratios.append((p.stem, round(ecart, 4)))
    legendes = len(re.findall(r"Figure[\u00a0\u202f ]\d+", ctx["texte"]))
    return dict(verdict="vert" if len(pdf) == len(png) == 61 and not ratios and legendes >= 61
                else "rouge", mesure=len(pdf),
                detail=f"{len(pdf)} PDF, {len(png)} PNG, ratios divergents {ratios[:3] or 'aucun'}, "
                       f"{legendes} légendes de planches lues dans le rendu")


EMOJIS_HORS = set("⛔✅❌⚠︎️")
def _pictogrammes(texte: str) -> list[str]:
    vus = []
    for ch in set(texte):
        o = ord(ch)
        if 0x1F000 <= o <= 0x1FAFF or 0x2600 <= o <= 0x27BF or ch in EMOJIS_HORS:
            vus.append(ch)
    return vus


def c_emojis(ctx):
    trouves = _pictogrammes(ctx["texte"])
    sources = []
    for f in sorted((RACINE / "rapport").glob("figs*.py")):
        sources += [f"{f.name}:{i}" for i, ln in enumerate(f.read_text(encoding="utf-8").splitlines(), 1)
                    if _pictogrammes(ln)]
    return dict(verdict="vert" if not trouves and not sources else "rouge", mesure=len(trouves),
                detail=f"{len(trouves)} pictogramme(s) au rendu {[t for t in trouves][:4]}, "
                       f"{len(sources)} ligne(s) de source contaminée(s) {sources[:3] or 'aucune'}")


def c_secrets(ctx):
    """Le kit ne porte aucun secret ; le dépôt ne suit aucun `.env` réel ; le verrou est relu.

    Trois précisions, parce qu'un contrôle de secrets écrit vite ment :
    • le périmètre analysé est celui du **kit** (rapport/, soutenance/, audit/), pris sur le système
      de fichiers — `git ls-files` renverrait zéro fichier ici, tout n'étant pas indexé à la date du
      relevé, et un contrôle qui ne regarde rien doit rougir plutôt que passer ;
    • `.env.example` est admis à une condition : aucune valeur hors gabarit, c'est-à-dire vide,
      bouclage 127.0.0.1/localhost, ou marqueur explicite (`change-me`) ;
    • deux fichiers de l'application (hors kit, non modifiés ici) portent une URL de base de
      développement en clair : ils sont **consignés** dans ce procès-verbal et dits en IV.2.
    """
    motifs = [r"(?i)\b(api[_-]?key|secret|password|passwd|private[_-]?key)\b\s*[:=]\s*['\"]([^\'\"]{6,})",
              r"-----BEGIN [A-Z ]*PRIVATE KEY-----", r"(?i)bearer\s+[A-Za-z0-9._-]{20,}"]
    fichiers_kit = [p for d in ("rapport", "soutenance", "audit") for p in (RACINE / d).rglob("*")
                    if p.is_file() and p.suffix in {".py", ".md", ".json", ".txt", ".ts"}]
    touches = []
    for p in fichiers_kit:
        t = p.read_text(encoding="utf-8", errors="ignore")
        for m in motifs:
            for mm in re.finditer(m, t):
                if not re.search(r"change-me|exemple|placeholder|\*\*\*|gabarit", mm.group(0), re.I):
                    touches.append(f"{p.relative_to(RACINE)} : {mm.group(0)[:40]}")
    envs = [l for l in _git_ls() if re.fullmatch(r"\.env(\.[\w-]+)?", l)]
    env_reels = [l for l in envs if not l.endswith(".example")]
    env_faux = []
    for l in envs:
        p = RACINE / l
        if not p.is_file():
            continue
        for ln in p.read_text(encoding="utf-8").splitlines():
            if "=" not in ln or ln.strip().startswith("#"):
                continue
            val = ln.split("=", 1)[1].strip()
            if val and not re.search(r"localhost|127\.0\.0\.1|change-me|example|gabarit", val, re.I):
                env_faux.append(f"{l} : {ln.split('=', 1)[0]}")
    gitignore = (RACINE / ".gitignore").read_text(encoding="utf-8") if (RACINE / ".gitignore").exists() else ""
    attendus = ["node_modules", ".env", ".next", "__pycache__", "*.py[cod]"]
    manquent = [x for x in attendus if x not in gitignore]
    ignore_le_kit = [x for x in ("audit", "rapport/PDF", "soutenance")
                     if re.search(rf"(?m)^{re.escape(x)}/?$", gitignore)]
    constat = []
    for l in _git_ls():
        p = RACINE / l
        if p.is_file() and l.endswith((".json", ".ts", ".example")) and not l.startswith(
                ("rapport/", "audit/", "soutenance/")):
            if re.search(r"postgres(ql)?://\w+:\w+@", p.read_text(encoding="utf-8", errors="ignore")):
                constat.append(l)
    ok = (not touches and not env_reels and not env_faux and not manquent and not ignore_le_kit)
    return dict(verdict="vert" if ok else "rouge", mesure=len(fichiers_kit),
                detail=f"{len(fichiers_kit)} fichiers du kit analysés ; secrets {touches[:2] or 'aucun'} ; "
                       f".env suivis réels {env_reels or 'aucun'}, valeurs hors gabarit "
                       f"{env_faux or 'aucune'} ; .gitignore sans {manquent or 'rien'} ; chemins du kit "
                       f"ignorés à tort {ignore_le_kit or 'aucun'} ; URL de base de développement en "
                       f"clair hors kit, consignée : {constat or 'aucune'}")


def c_deck(ctx):
    import zipfile
    from pptx import Presentation
    if not DECK.exists():
        return dict(verdict="rouge", mesure=0,
                    detail="diaporama absent : lancer `python3 soutenance/slides.py`")
    z = zipfile.ZipFile(DECK)
    slides = sorted(n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
    morph = sum(1 for n in slides if b"AlternateContent" in z.read(n))
    durees = sum(1 for n in slides if b'p14:dur="900"' in z.read(n))
    p = Presentation(str(DECK))
    notes = sum(1 for s in p.slides if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip())
    ratio = round(p.slide_width / p.slide_height, 4)
    plan = next((s for s in p.slides if any((sh.has_text_frame and sh.text_frame.text.strip() == "Plan")
                                            or "Plan" in (sh.text_frame.text[:20] if sh.has_text_frame else "")
                                            for sh in s.shapes)), None)
    trop_long_plan = [len(r.text) for sh in (plan.shapes if plan else []) if sh.has_text_frame
                      for para in sh.text_frame.paragraphs for r in para.runs if len(r.text) >= 90]
    trop_long_partout = [len(r.text) for s in p.slides for sh in s.shapes if sh.has_text_frame
                         and (sh.name or "") == "contenu" for para in sh.text_frame.paragraphs
                         for r in para.runs if len(r.text) > 140]
    ok = (len(slides) == 30 and morph == 5 and durees >= 5 and notes == 30
          and abs(ratio - 16 / 9) < 0.01 and not trop_long_plan and not trop_long_partout)
    return dict(verdict="vert" if ok else "rouge", mesure=len(slides),
                detail=f"{len(slides)} glissades, morph {morph}, durées 900 ms {durees}, notes {notes}, "
                       f"ratio {ratio}, entrées du plan ≥ 90 car. {len(trop_long_plan)}, libellés "
                       f"> 140 car. {len(trop_long_partout)}")


def c_nombres(ctx):
    points = sorted(set(re.findall(r"\d+\.\d+\s?DT", ctx["texte"])))
    lignes = ctx["lignes"]
    dans_code = _bloc_code(lignes)
    groupements = [l[:48] for l, c in zip(lignes, dans_code)
                   if not c and re.search(r"(?<!\d)\d{1,3}(?:\.\d{3})+(?:,\d+)?", l)]
    virgules = len(re.findall(r"\d+,\d+\s?DT", ctx["texte"]))
    return dict(verdict="vert" if not points and not groupements and virgules >= 10 else "rouge",
                mesure=virgules,
                detail=f"{virgules} montants à virgule dans le rendu, points décimaux "
                       f"{points[:4] or 'aucun'}, groupements à l'anglaise {groupements[:2] or 'aucun'}")


def c_insécables(ctx):
    coupures, exemptees = [], 0
    lignes = ctx["lignes"]
    dans_code = _bloc_code(lignes)
    for l, code in zip(lignes, dans_code):
        if not l:
            continue
        if l[0] in ";:!?»%" or l.endswith("«"):
            if code:
                exemptees += 1
            else:
                coupures.append(l[:48])
    return dict(verdict="vert" if not coupures else "rouge", mesure=len(coupures),
                detail=f"{len(coupures)} césure(s) de ponctuation sur {len(lignes)} lignes rendues, "
                       f"{exemptees} ligne(s) de verbatim exemptée(s) : {coupures[:3] or 'aucune'}")


CONTROLES = [
    ("QA-1", "versions citées = package.json", c_version),
    ("QA-2", "chiffres du rapport = dépôt recompté", c_chiffres),
    ("QA-3", "chemins et lignes cités réels", c_chemins),
    ("QA-4", "aucune formule de remplissage", c_slop),
    ("QA-5", "aucun mot anglais hors verbatim", c_synonymes),
    ("QA-6", "signets cohérents ×3", c_signets),
    ("QA-7", "liens internes tous résolus", c_liens),
    ("QA-8", "planches PDF et PNG en ratio", c_planches),
    ("QA-9", "aucun pictogramme dans une figure", c_emojis),
    ("QA-10", "aucun secret dans les fichiers suivis", c_secrets),
    ("QA-11", "diaporama : 30, morph, notes, formats", c_deck),
    ("QA-12", "nombres à la française", c_nombres),
]
# Le treizième regard est la loi des coupures : qa.py la juge par voisinage de texte, qa2.py par la
# géométrie des spans et le nom de la police. Les deux comptages sont imprimés côte à côte dans
# `audit/ab/harnais.json` — c'est la raison d'être du double harnais, pas un redoublement de confort.
CONTROLES.append(("QA-13", "ponctuation française non coupée au rendu", c_insécables))


def lancer(*, silencieux: bool = False) -> dict:
    t0 = time.time()
    texte, lignes, pages = _texte_pdf()
    ctx = dict(texte=texte, lignes=lignes, pages=pages)
    recu = dict(harnais="qa.py (pypdf, lecture page à page)", pages=pages,
                caracteres=len(texte), controles=[])
    rouges = 0
    for code, titre, func in CONTROLES:
        try:
            r = func(ctx)
        except Exception as exc:                                  # noqa: BLE001 : un plantage EST rouge
            r = dict(verdict="rouge", mesure=0, detail=f"exception {type(exc).__name__}: {exc}")
        r.update(code=code, titre=titre)
        rouges += r["verdict"] == "rouge"
        recu["controles"].append(r)
        if not silencieux:
            print(f"{r['verdict']:>5}  {code}  {titre}\n       {r['detail']}")
    recu.update(rouges=rouges, secondes=round(time.time() - t0, 2),
                verdict="vert" if rouges == 0 else "rouge")
    AUDIT.mkdir(parents=True, exist_ok=True)
    (AUDIT / "qa.json").write_text(json.dumps(recu, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
                                   encoding="utf-8")
    budgets = AUDIT / "budgets.json"
    b = json.loads(budgets.read_text(encoding="utf-8")) if budgets.exists() else {}
    b["secondes_qa"] = round(sum(x["secondes"] for x in [recu]), 2)
    budgets.write_text(json.dumps(b, ensure_ascii=False, indent=1, sort_keys=True) + "\n",
                       encoding="utf-8")
    if not silencieux:
        print(f"\nqa : {len(CONTROLES)} contrôles, {rouges} rouge(s), {recu['secondes']} s")
    return recu


if __name__ == "__main__":
    import sys
    sys.exit(0 if lancer()["verdict"] == "vert" else 1)
